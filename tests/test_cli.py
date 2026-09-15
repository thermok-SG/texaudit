from pathlib import Path

from texaudit.audit import audit
from texaudit.parser import parse_manuscript
from texaudit.profiles import load_profile


def test_parse_sample():
    fixture = Path(__file__).parent / "fixtures" / "sample.tex"
    stats = parse_manuscript(fixture)
    assert stats.abstract_words > 0
    assert stats.plain_language_summary_words > 0
    assert stats.figure_count == 1
    assert stats.table_count == 1
    assert stats.equation_count == 1
    assert stats.key_point_count == 2
    assert stats.body_words > 0
    assert any("child.tex" in item for item in stats.files_read)


def test_parse_official_agu_2025_structure():
    fixture = Path(__file__).parent / "fixtures" / "agu2025.tex"
    stats = parse_manuscript(fixture)
    assert stats.abstract_words == 6
    assert stats.plain_language_summary_words == 7
    assert stats.key_point_count == 2
    assert stats.key_point_max_characters == len("Second formatted key point")
    assert stats.body_words == 8
    assert stats.agu_word_count == 14


def test_agu_profile_audits_sample():
    fixture = Path(__file__).parent / "fixtures" / "sample.tex"
    profile, path = load_profile("agu-jgr-earth-surface")
    report = audit(parse_manuscript(fixture), profile, path)
    names = {check.name for check in report.checks}
    assert "Publication units" in names
    assert "Abstract" in names


def test_common_agu_profiles_are_available():
    from texaudit.profiles import builtin_profile_names

    names = builtin_profile_names()
    assert "agu-grl" in names
    assert "agu-jgr-earth-surface" in names
    assert "agu-jgr-solid-earth" in names
    assert "agu-tectonics" in names
    assert "agu-water-resources-research" in names
    assert "_agu-research-article" not in names


def test_grl_profile_overrides_shared_agu_rules():
    from texaudit.models import ManuscriptStats

    profile, path = load_profile("agu-grl")
    assert profile["sections"]["abstract"]["max_words"] == 149
    assert profile["sections"]["plain_language_summary"]["required"] is True
    stats = ManuscriptStats(source="example.tex", abstract_words=100, body_words=6001)
    report = audit(stats, profile, path)
    publication_units = next(check for check in report.checks if check.name == "Publication units")
    assert publication_units.status == "FAIL"
    assert publication_units.expected == 12


def test_required_key_points_have_a_minimum():
    from texaudit.models import ManuscriptStats

    profile, path = load_profile("agu-jgr-earth-surface")
    report = audit(ManuscriptStats(source="example.tex", abstract_words=10), profile, path)
    key_points = next(check for check in report.checks if check.name == "Key points")
    assert key_points.status == "FAIL"
    assert key_points.expected == "1-3"


def test_fully_open_access_agu_profile_reports_pus_without_fee_warning():
    from texaudit.models import ManuscriptStats

    profile, path = load_profile("agu-water-resources-research")
    report = audit(ManuscriptStats(source="example.tex", body_words=13000), profile, path)
    publication_units = next(check for check in report.checks if check.name == "Publication units")
    assert publication_units.actual == 26
    assert publication_units.status == "INFO"
    assert publication_units.expected is None


def test_parse_docx(tmp_path):
    from docx import Document

    path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_heading("Abstract", level=1)
    doc.add_paragraph("This is a compact abstract with enough words to count.")
    doc.add_heading("Plain Language Summary", level=1)
    doc.add_paragraph("This explains the study in accessible language.")
    doc.add_heading("Key Points", level=1)
    doc.add_paragraph("First point", style="List Bullet")
    doc.add_paragraph("Second point", style="List Bullet")
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("This is the manuscript body text.")
    doc.add_paragraph("Figure 1. A useful figure caption.", style="Caption")
    doc.add_paragraph("Table 1. A useful table caption.", style="Caption")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Header one"
    table.cell(0, 1).text = "Header two"
    doc.add_heading("References", level=1)
    doc.add_paragraph("Smith, A. Example reference.")
    doc.save(path)

    from texaudit.docx_reader import parse_docx
    stats = parse_docx(path)
    assert stats.abstract_words > 0
    assert stats.plain_language_summary_words > 0
    assert stats.key_point_count == 2
    assert stats.figure_caption_words > 0
    assert stats.table_caption_words > 0
    assert stats.table_count == 1
    assert stats.reference_words > 0


def test_parse_agu_word_template_styles(tmp_path):
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from texaudit.docx_reader import parse_docx

    path = tmp_path / "agu-template.docx"
    doc = Document()
    for name in ("Heading-Main", "Heading-Secondary", "Affiliation", "Note", "Key Points", "Abstract", "Figure or Table Caption"):
        doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("A concise manuscript title", style="Title")
    doc.add_paragraph("Author Name", style="Affiliation")
    doc.add_paragraph("Corresponding author details", style="Note")
    doc.add_paragraph("Key Points", style="Heading-Main")
    doc.add_paragraph("First point", style="Key Points")
    doc.add_paragraph("Abstract", style="Heading-Main")
    doc.add_paragraph("Short abstract text", style="Abstract")
    doc.add_paragraph("Plain Language Summary", style="Abstract")
    doc.add_paragraph("Short plain summary", style="Abstract")
    doc.add_paragraph("1 Introduction", style="Heading-Main")
    doc.add_paragraph("Four words of body text")
    doc.add_paragraph("2 Materials and Methods", style="Heading-Main")
    doc.add_paragraph("2.1 Tests", style="Heading-Secondary")
    doc.add_paragraph("Three method words")
    doc.add_paragraph("3 Results", style="Heading-Main")
    doc.add_paragraph("Two result words")
    caption = "Figure 1. One numbered caption"
    doc.add_paragraph(caption)
    doc.add_paragraph(caption, style="Figure or Table Caption")
    doc.add_paragraph("Table 1. One table caption", style="Figure or Table Caption")
    doc.add_paragraph("References")
    doc.add_paragraph("First reference")
    doc.add_paragraph("Second reference")
    doc.save(path)

    stats = parse_docx(path)
    assert stats.title_words == 4
    assert stats.body_words == 11
    assert stats.methods_words == 3
    assert stats.figure_count == 1
    assert stats.table_count == 1
    assert stats.figure_caption_words == 5
    assert stats.reference_count == 2


def test_nature_geoscience_limits_and_guidance():
    from texaudit.models import ManuscriptStats

    profile, path = load_profile("nature-geoscience-article")
    stats = ManuscriptStats(
        source="example.docx",
        title_characters=91,
        abstract_words=150,
        body_words=4000,
        methods_words=1500,
        figure_count=5,
        table_count=2,
        reference_count=51,
        section_names=["data availability"],
    )
    report = audit(stats, profile, path)
    checks = {check.name: check for check in report.checks}
    assert checks["Title"].status == "FAIL"
    assert checks["Main text"].actual == 2500
    assert checks["Main text"].status == "PASS"
    assert checks["Display items"].status == "FAIL"
    assert checks["References"].status == "WARNING"


def test_elsevier_highlight_limits():
    from texaudit.models import ManuscriptStats

    profile, path = load_profile("elsevier-geomorphology")
    stats = ManuscriptStats(source="highlights.docx", highlight_count=3, highlight_max_characters=86)
    report = audit(stats, profile, path)
    checks = {check.name: check for check in report.checks}
    assert checks["Highlights"].status == "PASS"
    assert checks["Longest highlight"].status == "FAIL"
