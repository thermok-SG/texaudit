from __future__ import annotations

import re
from pathlib import Path

from docx import Document

from .counting import count_characters, count_words
from .models import ManuscriptStats
from .parser import ACK_NAMES, APPENDIX_NAMES, OPEN_RESEARCH_NAMES, normalise_name

ABSTRACT_NAMES = {"abstract"}
PLS_NAMES = {"plain language summary", "plain-language summary", "plainlanguagesummary", "pls"}
KEYPOINT_NAMES = {"key points", "keypoints", "key_points"}
HIGHLIGHT_NAMES = {"highlights", "research highlights"}
REFERENCE_NAMES = {"references", "bibliography", "reference list"}
CAPTION_RE = re.compile(r"^\s*(figure|fig\.?|table)\s*(?:s\.?\s*)?(\d+[a-zA-Z]?)?\s*[:.\-]", re.I)


def _paragraph_is_heading(paragraph) -> bool:
    style_name = (getattr(paragraph.style, "name", "") or "").lower()
    return style_name.startswith("heading") or style_name in {"title", "subtitle"}


def _heading_level(paragraph) -> int | None:
    style_name = (getattr(paragraph.style, "name", "") or "").lower()
    match = re.match(r"heading\s*[- ]?\s*(\d+)", style_name)
    if match:
        return int(match.group(1))
    if style_name in {"heading-main", "heading main"}:
        return 1
    if style_name in {"heading-secondary", "heading secondary"}:
        return 2
    return None


def _section_name(value: str) -> str:
    return re.sub(r"^\d+(?:\.\d+)*\s+", "", value)


def _caption_kind(paragraph) -> str | None:
    text = paragraph.text.strip()
    style_name = (getattr(paragraph.style, "name", "") or "").lower()
    match = CAPTION_RE.match(text)
    if "caption" in style_name or match:
        if match and match.group(1).lower().startswith("t"):
            return "table"
        if match and match.group(1).lower().startswith("f"):
            return "figure"
        # Word's Caption style does not identify item type; infer from text where possible.
        if text.lower().startswith("table"):
            return "table"
        return "figure"
    return None


def _drawing_count(document) -> int:
    # Inline shapes miss floating images; counting w:drawing catches both.
    try:
        return len(document.element.body.xpath(".//w:drawing"))
    except Exception:
        return len(document.inline_shapes)


def _equation_count(document) -> int:
    try:
        displayed = document.element.body.xpath(".//m:oMathPara")
        if displayed:
            return len(displayed)
        return len(document.element.body.xpath(".//m:oMath"))
    except Exception:
        return 0


def _citation_count(document) -> int:
    try:
        instructions = document.element.body.xpath(".//w:instrText")
        return sum(1 for item in instructions if "CITATION" in (item.text or "").upper() or " CITE " in f" {(item.text or '').upper()} ")
    except Exception:
        return 0


def parse_docx(path: Path) -> ManuscriptStats:
    path = path.expanduser().resolve()
    if not path.exists():
        raise FileNotFoundError(f"DOCX file not found: {path}")
    document = Document(path)
    stats = ManuscriptStats(source=str(path), files_read=[str(path)])
    stats.figure_count = _drawing_count(document)
    stats.table_count = len(document.tables)
    stats.equation_count = _equation_count(document)
    stats.citation_command_count = _citation_count(document)

    current = "body"
    key_items: list[str] = []
    highlights: list[str] = []
    seen_heading_names: list[str] = []
    content_started = False
    seen_captions: set[tuple[str, str]] = set()
    figure_caption_numbers: set[str] = set()
    table_caption_numbers: set[str] = set()

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        name = normalise_name(text.rstrip(":"))
        section_name = _section_name(name)
        style_name = (getattr(paragraph.style, "name", "") or "").lower()
        if style_name == "title":
            if section_name in HIGHLIGHT_NAMES:
                current = section_name
                content_started = True
                seen_heading_names.append(section_name)
                continue
            if stats.title_words == 0:
                stats.title_words = count_words(text)
                stats.title_characters = count_characters(text)
            continue
        is_named_heading = section_name in (ABSTRACT_NAMES | PLS_NAMES | KEYPOINT_NAMES | HIGHLIGHT_NAMES | REFERENCE_NAMES | ACK_NAMES | APPENDIX_NAMES | OPEN_RESEARCH_NAMES)
        if _paragraph_is_heading(paragraph) or is_named_heading:
            content_started = True
            if is_named_heading:
                current = section_name
                seen_heading_names.append(section_name)
                continue
            # Styled headings are retained as body structure but do not count as prose.
            seen_heading_names.append(name)
            level = _heading_level(paragraph)
            if level == 1 or current != "methods":
                current = "methods" if section_name in {"methods", "online methods", "materials and methods"} else ("appendix" if section_name.startswith("appendix") else "body")
            continue

        if not content_started and style_name in {"affiliation", "author", "authors", "note", "subtitle"}:
            continue

        caption_kind = _caption_kind(paragraph)
        if caption_kind:
            match = CAPTION_RE.match(text)
            number = match.group(2).lower() if match and match.group(2) else ""
            identity = (caption_kind, number or normalise_name(text))
            if identity not in seen_captions:
                seen_captions.add(identity)
                if caption_kind == "figure":
                    stats.figure_caption_words += count_words(text)
                    if number:
                        figure_caption_numbers.add(number)
                else:
                    stats.table_caption_words += count_words(text)
                    if number:
                        table_caption_numbers.add(number)
            continue

        if current in ABSTRACT_NAMES:
            stats.abstract_words += count_words(text)
        elif current in PLS_NAMES:
            stats.plain_language_summary_words += count_words(text)
        elif current in KEYPOINT_NAMES:
            key_items.append(text)
        elif current in HIGHLIGHT_NAMES:
            highlights.append(text.lstrip("-•* ").strip())
        elif current in REFERENCE_NAMES:
            stats.reference_words += count_words(text)
            stats.reference_count += 1
        elif current in ACK_NAMES:
            stats.acknowledgements_words += count_words(text)
        elif current in APPENDIX_NAMES or current.startswith("appendix"):
            stats.appendix_words += count_words(text)
        elif current in OPEN_RESEARCH_NAMES:
            stats.open_research_words += count_words(text)
        else:
            stats.body_words += count_words(text)
            if current == "methods":
                stats.methods_words += count_words(text)

    stats.section_names = seen_heading_names
    stats.key_point_count = len(key_items)
    stats.key_point_max_characters = max((count_characters(item) for item in key_items), default=0)
    stats.highlight_count = len(highlights)
    stats.highlight_max_characters = max((count_characters(item) for item in highlights), default=0)
    stats.figure_count = max(stats.figure_count, len(figure_caption_numbers))
    stats.table_count = max(stats.table_count, len(table_caption_numbers))

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                stats.table_text_words += count_words(cell.text)

    # Normal Word text cannot expose TeX-style unresolved refs, but literal markers are still useful.
    all_text = "\n".join(p.text for p in document.paragraphs)
    stats.unresolved_marker_count = all_text.count("??")
    if stats.figure_count == 0:
        stats.warnings.append("No embedded Word drawings detected. Linked images or externally supplied figures are not counted.")
    stats.warnings.append("DOCX citation detection only recognises Word field-code citations; Zotero, EndNote, and plain-text citations may not be counted.")
    return stats
